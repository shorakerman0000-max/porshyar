# -*- coding: utf-8 -*-
"""پرسش‌یار آفلاین با Flet؛ بخش‌ها و سؤال‌های دلخواه، ذخیره SQLite و خروجی Excel."""
import json, sqlite3, os, sys
from datetime import datetime
from pathlib import Path
import flet as ft

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None

APP_DIR = Path.home() / ".porshyar"
APP_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = APP_DIR / "porshyar.sqlite3"
EXPORT_DIR = APP_DIR / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

TEMPLATES = [
("اطلاعات عمومی و مشخصات فردی", ["سن", "میزان تحصیلات", "وضعیت اشتغال", "شهر محل سکونت"]),
("ارزیابی آموزشی و تحصیلی", ["موضوع یا مقطع آموزشی", "میزان رضایت از کیفیت آموزش", "مهم‌ترین نیاز آموزشی شما چیست؟"]),
("رضایت‌سنجی و نظرسنجی", ["از کدام خدمت استفاده کرده‌اید؟", "میزان رضایت کلی شما", "پیشنهاد شما چیست؟"]),
("پرسشنامه پژوهشی و تحقیقاتی", ["عنوان یا موضوع پژوهش", "نظر یا تجربه شما درباره موضوع چیست؟"]),
("ارزیابی عملکرد و محیط کار", ["واحد یا نوع فعالیت", "مهم‌ترین چالش کاری شما چیست؟"]),
("نیازسنجی اجتماعی و خدماتی", ["منطقه یا جامعه مورد بررسی", "مهم‌ترین نیاز فعلی شما چیست؟"]),
("ارزیابی برنامه‌ها و رویدادها", ["نام برنامه یا رویداد", "میزان رضایت کلی از برنامه", "چه پیشنهادی دارید؟"]),
]
TEMPLATE_BOXES = [
("اطلاعات واردشده را تأیید می‌کنم.", "آیا مایلید برای پیگیری با شما تماس گرفته شود؟"),
("با ثبت پاسخ‌ها موافقم.", "آیا این دوره را به دیگران پیشنهاد می‌کنید؟"),
("پاسخ‌ها بازتاب‌دهندهٔ نظر واقعی من است.", "آیا از خدمت دریافت‌شده رضایت دارید؟"),
("با استفاده از پاسخ‌ها در چارچوب پژوهش موافقم.", "آیا قبلاً در پژوهشی مشابه شرکت کرده‌اید؟"),
("پاسخ‌ها را تا حد امکان دقیق وارد کرده‌ام.", "آیا ابزار کافی برای انجام وظایف دارید؟"),
("اطلاعات بالا را تا حد امکان دقیق وارد کرده‌ام.", "آیا به خدمات موردنظر دسترسی دارید؟"),
("در این نظرسنجی شرکت می‌کنم.", "آیا در برنامهٔ مشابه بعدی شرکت خواهید کرد؟"),
]

def question(kind="essay", prompt="", label="بله، تأیید می‌کنم", yes="بله", no="خیر"):
    return {"type": kind, "prompt": prompt, "label": label, "yes_label": yes, "no_label": no}

def normalize(raw):
    try:
        data = json.loads(raw or "{}")
    except Exception:
        data = {}
    if isinstance(data, dict) and isinstance(data.get("sections"), list):
        sections = []
        for i, sec in enumerate(data["sections"], 1):
            qs = []
            for it in sec.get("questions", []):
                typ = it.get("type", "essay")
                if typ not in ("essay", "checkbox", "yesno"):
                    typ = "essay"
                qs.append(question(typ, str(it.get("prompt", "")), str(it.get("label", "بله، تأیید می‌کنم")),
                                   str(it.get("yes_label", "بله")), str(it.get("no_label", "خیر"))))
            sections.append({"title": str(sec.get("title") or f"بخش {i}"), "questions": qs})
        return {"sections": sections or [{"title": "بخش ۱", "questions": []}]}
    # پشتیبانی از ساختار نسخه‌های قبلی
    old = data if isinstance(data, dict) else {"questions": data if isinstance(data, list) else []}
    qs = [question("essay", str(x)) for x in old.get("questions", []) if str(x).strip()]
    if str(old.get("checkbox_question", "")).strip():
        qs.append(question("checkbox", str(old["checkbox_question"]), str(old.get("checkbox_label", "بله، تأیید می‌کنم"))))
    if str(old.get("yesno_question", "")).strip():
        qs.append(question("yesno", str(old["yesno_question"]), yes=str(old.get("yes_label", "بله")), no=str(old.get("no_label", "خیر"))))
    return {"sections": [{"title": "بخش سؤال‌ها", "questions": qs}]}

class Store:
    def connect(self):
        c = sqlite3.connect(DB_PATH)
        c.row_factory = sqlite3.Row
        return c

    def __init__(self):
        with self.connect() as c:
            c.execute("""CREATE TABLE IF NOT EXISTS questionnaires(
                id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL,
                questions TEXT NOT NULL, is_template INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL)""")
            c.execute("""CREATE TABLE IF NOT EXISTS responses(
                id INTEGER PRIMARY KEY AUTOINCREMENT, questionnaire_id INTEGER NOT NULL,
                questionnaire_title TEXT NOT NULL, answers TEXT NOT NULL, created_at TEXT NOT NULL)""")
            count = c.execute("SELECT COUNT(*) FROM questionnaires WHERE is_template=1").fetchone()[0]
            now = datetime.now().isoformat(timespec="seconds")
            if count == 0:
                for i, (title, essays) in enumerate(TEMPLATES):
                    box, yn = TEMPLATE_BOXES[i]
                    cfg = {"sections": [
                        {"title": "سؤال‌های تشریحی", "questions": [question("essay", x) for x in essays]},
                        {"title": "سؤال تیک‌زدنی", "questions": [question("checkbox", box)]},
                        {"title": "سؤال بله/خیر", "questions": [question("yesno", yn)]},
                    ]}
                    c.execute("INSERT INTO questionnaires(title,questions,is_template,created_at) VALUES(?,?,1,?)",
                              (title, json.dumps(cfg, ensure_ascii=False), now))
            else:
                # قالب‌های قدیمی را بدون حذف پاسخ‌ها به ساختار بخش‌بندی‌شده ارتقا می‌دهد.
                for i, (title, essays) in enumerate(TEMPLATES):
                    row = c.execute("SELECT id,questions FROM questionnaires WHERE title=? AND is_template=1", (title,)).fetchone()
                    if row:
                        cfg = normalize(row["questions"])
                        box, yn = TEMPLATE_BOXES[i]
                        # قالب‌های موجود را هم ارتقا بده؛ سؤال‌های فعلی کاربر را نگه دار
                        # و اگر نوع تیک‌زدنی یا بله/خیر ندارند، بخش مربوط را اضافه کن.
                        all_questions = [q for sec in cfg["sections"] for q in sec["questions"]]
                        if not all_questions:
                            cfg["sections"][0]["questions"] = [question("essay", x) for x in essays]
                            all_questions = cfg["sections"][0]["questions"]
                        if not any(q["type"] == "checkbox" for q in all_questions):
                            cfg["sections"].append({
                                "title": "سؤال تیک‌زدنی",
                                "questions": [question("checkbox", box)]
                            })
                        if not any(q["type"] == "yesno" for q in all_questions):
                            cfg["sections"].append({
                                "title": "سؤال بله/خیر",
                                "questions": [question("yesno", yn)]
                            })
                        c.execute("UPDATE questionnaires SET questions=? WHERE id=?",
                                  (json.dumps(cfg, ensure_ascii=False), row["id"]))

    def all_questionnaires(self):
        with self.connect() as c:
            return c.execute("SELECT * FROM questionnaires ORDER BY is_template DESC,id DESC").fetchall()

    def get_questionnaire(self, qid):
        with self.connect() as c:
            return c.execute("SELECT * FROM questionnaires WHERE id=?", (qid,)).fetchone()

    def save_questionnaire(self, qid, title, cfg):
        payload = json.dumps(cfg, ensure_ascii=False)
        with self.connect() as c:
            if qid:
                c.execute("UPDATE questionnaires SET title=?,questions=? WHERE id=?", (title, payload, qid))
            else:
                c.execute("INSERT INTO questionnaires(title,questions,is_template,created_at) VALUES(?,?,0,?)",
                          (title, payload, datetime.now().isoformat(timespec="seconds")))

    def save_response(self, qid, title, cfg, answers):
        with self.connect() as c:
            cur = c.execute("INSERT INTO responses(questionnaire_id,questionnaire_title,answers,created_at) VALUES(?,?,?,?)",
                (qid, title, json.dumps({"config": cfg, "answers": answers}, ensure_ascii=False),
                 datetime.now().isoformat(timespec="seconds")))
            return cur.lastrowid

    def responses(self, qid=None):
        with self.connect() as c:
            if qid is None:
                return c.execute("SELECT * FROM responses ORDER BY id DESC").fetchall()
            return c.execute("SELECT * FROM responses WHERE questionnaire_id=? ORDER BY id DESC", (qid,)).fetchall()

def main(page: ft.Page):
    store = Store()
    page.title = "پرسش‌یار"
    page.rtl = True
    page.padding = 16
    page.scroll = ft.ScrollMode.AUTO
    blue, green, purple, gray = "#1F5D8F", "#147D64", "#6850A4", "#667085"

    def notify(message):
        page.snack_bar = ft.SnackBar(content=ft.Text(str(message), text_align=ft.TextAlign.RIGHT))
        page.snack_bar.open = True
        page.update()

    def text(value, size=14, color=None, bold=False):
        return ft.Text(str(value), size=size, color=color,
                       weight=ft.FontWeight.BOLD if bold else ft.FontWeight.NORMAL,
                       text_align=ft.TextAlign.RIGHT)

    def button(label, callback, color=blue, icon=None):
        return ft.Button(label, on_click=callback, bgcolor=color, color="white", icon=icon)

    def field(label, value="", multiline=False):
        return ft.TextField(label=label, value=value, multiline=multiline,
            min_lines=2 if multiline else 1, max_lines=6 if multiline else 1,
            rtl=True, text_align=ft.TextAlign.RIGHT)

    def header(title, back):
        return ft.Row([ft.IconButton(ft.Icons.ARROW_BACK, on_click=back),
                       ft.Text(title, size=20, weight=ft.FontWeight.BOLD, expand=True, text_align=ft.TextAlign.RIGHT)])

    def home():
        page.controls.clear()
        controls = [text("پرسش‌یار", 30, blue, True),
            text("پرسشنامه‌ساز آفلاین با بخش‌ها و انواع سؤال دلخواه", 14, gray),
            ft.Row([button("پرسشنامه جدید", lambda e: editor(None), green, ft.Icons.ADD),
                    button("همه پاسخ‌ها", lambda e: responses_view(None), purple)], wrap=True),
            ft.Divider(), text("قالب‌ها و پرسشنامه‌ها", 17, bold=True)]
        for row in store.all_questionnaires():
            qid = row["id"]
            controls.append(ft.Card(content=ft.Container(padding=12, content=ft.Column([
                text(f'{row["title"]} · {"قالب آماده" if row["is_template"] else "پرسشنامه من"}', 15, bold=True),
                ft.Row([ft.OutlinedButton("تکمیل", on_click=lambda e, i=qid: answer_view(i)),
                        ft.OutlinedButton("ویرایش", on_click=lambda e, i=qid: editor(i)),
                        ft.OutlinedButton("پاسخ‌ها", on_click=lambda e, i=qid: responses_view(i))], wrap=True)
            ]))))
        page.add(ft.Column(controls, spacing=12, scroll=ft.ScrollMode.AUTO, expand=True))
        page.update()

    def editor(qid=None):
        row = store.get_questionnaire(qid) if qid else None
        cfg = normalize(row["questions"]) if row else {"sections": [{"title": "بخش ۱", "questions": [question("essay", "")]}]}
        title_field = field("عنوان پرسشنامه", row["title"] if row else "")
        main_col = None

        def render():
            if main_col is None:
                return
            cards = []
            for si, sec in enumerate(cfg["sections"]):
                titlebox = field("نام بخش", sec["title"])
                titlebox.on_change = lambda e, i=si, f=titlebox: cfg["sections"][i].update({"title": f.value or ""})
                qcards = []
                for qi, item in enumerate(sec["questions"]):
                    kind = ft.Dropdown(label="نوع سؤال", value=item["type"],
                        options=[ft.dropdown.Option("essay", "تشریحی"), ft.dropdown.Option("checkbox", "تیک‌زدنی"), ft.dropdown.Option("yesno", "بله/خیر")],
                        rtl=True)
                    prompt = field("متن سؤال", item["prompt"])
                    check_label = field("متن کنار تیک", item.get("label", "بله، تأیید می‌کنم"))
                    yes_label = field("متن گزینه بله", item.get("yes_label", "بله"))
                    no_label = field("متن گزینه خیر", item.get("no_label", "خیر"))

                    def sync(e=None, i=si, j=qi, k=kind, p=prompt, cl=check_label, yl=yes_label, nl=no_label):
                        it = cfg["sections"][i]["questions"][j]
                        it.update({"type": k.value or "essay", "prompt": p.value or "", "label": cl.value or "",
                                   "yes_label": yl.value or "بله", "no_label": nl.value or "خیر"})
                        update_visibility(k, cl, yl, nl)
                    kind.on_change = sync
                    prompt.on_change = sync
                    check_label.on_change = sync
                    yes_label.on_change = sync
                    no_label.on_change = sync

                    def update_visibility(k, cl, yl, nl):
                        cl.visible = k.value == "checkbox"
                        yl.visible = nl.visible = k.value == "yesno"

                    update_visibility(kind, check_label, yes_label, no_label)

                    def delete_question(e, i=si, j=qi):
                        cfg["sections"][i]["questions"].pop(j)
                        render()
                    qcards.append(ft.Container(padding=10, border=ft.Border.all(1, "#D0D5DD"), border_radius=8,
                        content=ft.Column([ft.Row([text(f"سؤال {qi+1}", 14, bold=True),
                            ft.TextButton("حذف سؤال", on_click=delete_question)], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                            kind, prompt, check_label, yes_label, no_label], spacing=7)))

                def add_question(e, i=si):
                    cfg["sections"][i]["questions"].append(question("essay", ""))
                    render()

                def delete_section(e, i=si):
                    if len(cfg["sections"]) <= 1:
                        notify("حداقل یک بخش باید باقی بماند.")
                        return
                    cfg["sections"].pop(i)
                    render()

                cards.append(ft.Card(content=ft.Container(padding=12, content=ft.Column([
                    ft.Row([text(f"بخش {si+1}", 16, blue, True),
                            ft.TextButton("حذف بخش", on_click=delete_section)], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                    titlebox, *qcards, button("افزودن سؤال به این بخش", add_question, blue, ft.Icons.ADD)
                ], spacing=10))))
            main_col.controls = [header("ساخت / ویرایش پرسشنامه", lambda e: home()), title_field,
                text("هر بخش می‌تواند هر تعداد سؤال تشریحی، تیک‌زدنی یا بله/خیر داشته باشد.", 14, gray),
                *cards, button("افزودن بخش جدید", add_section, purple, ft.Icons.ADD),
                button("ذخیره پرسشنامه", save, green, ft.Icons.SAVE)]
            page.update()

        def save(e):
            title = (title_field.value or "").strip()
            if not title:
                notify("عنوان پرسشنامه را وارد کن.")
                return
            for sec in cfg["sections"]:
                sec["title"] = (sec.get("title") or "").strip() or "بخش بدون عنوان"
                sec["questions"] = [it for it in sec["questions"] if (it.get("prompt") or "").strip()]
            if not any(sec["questions"] for sec in cfg["sections"]):
                notify("حداقل یک سؤال با متن وارد کن.")
                return
            try:
                store.save_questionnaire(qid, title, cfg)
            except Exception as ex:
                notify(f"ذخیره انجام نشد: {ex}")
                return
            home()
            notify("پرسشنامه ذخیره شد.")

        def add_section(e):
            cfg["sections"].append({"title": f"بخش {len(cfg['sections'])+1}", "questions": [question("essay", "")]})
            render()

        main_col = ft.Column([], spacing=12, scroll=ft.ScrollMode.AUTO, expand=True)
        page.controls.clear()
        page.add(main_col)
        render()

    def answer_view(qid):
        row = store.get_questionnaire(qid)
        if not row:
            notify("پرسشنامه پیدا نشد.")
            return
        cfg = normalize(row["questions"])
        skip_controls, answer_controls = [], {}
        section_cards = []
        for si, sec in enumerate(cfg["sections"]):
            skip = ft.Checkbox(label=f"از بخش «{sec['title']}» صرف‌نظر می‌کنم", value=False)
            skip_controls.append(skip)
            q_controls = []
            for qi, item in enumerate(sec["questions"]):
                key = (si, qi)
                if item["type"] == "essay":
                    control = field(item["prompt"], "", True)
                elif item["type"] == "checkbox":
                    control = ft.Checkbox(label=f'{item["prompt"]}\n{item.get("label", "بله، تأیید می‌کنم")}', value=False)
                else:
                    control = ft.RadioGroup(content=ft.Column([
                        ft.Radio(value="yes", label=item.get("yes_label", "بله")),
                        ft.Radio(value="no", label=item.get("no_label", "خیر"))]), value=None)
                answer_controls[key] = control
                q_controls.append(control)
            content = ft.Column(q_controls, spacing=8)
            def on_skip(e, c=content, s=skip):
                c.visible = not bool(s.value)
                page.update()
            skip.on_change = on_skip
            section_cards.append(ft.Card(content=ft.Container(padding=12, content=ft.Column([
                text(sec["title"], 17, blue, True), skip, content], spacing=8))))

        def save_answer(e):
            answers = {}
            for si, sec in enumerate(cfg["sections"]):
                skipped = bool(skip_controls[si].value)
                items = {}
                for qi, item in enumerate(sec["questions"]):
                    control = answer_controls[(si, qi)]
                    if skipped:
                        value = None
                    elif item["type"] == "essay":
                        value = control.value or ""
                    elif item["type"] == "checkbox":
                        value = bool(control.value)
                    else:
                        value = control.value
                    items[str(qi)] = value
                answers[str(si)] = {"skipped": skipped, "items": items}
            try:
                rid = store.save_response(qid, row["title"], cfg, answers)
            except Exception as ex:
                notify(f"ذخیره پاسخ ناموفق بود: {ex}")
                return
            home()
            notify(f"پاسخ با شناسهٔ {rid} ذخیره شد.")

        page.controls.clear()
        page.add(ft.Column([header("تکمیل پرسشنامه", lambda e: home()), text(row["title"], 21, blue, True),
                            *section_cards, button("ذخیره پاسخ", save_answer, green, ft.Icons.SAVE)],
                           spacing=12, scroll=ft.ScrollMode.AUTO, expand=True))
        page.update()

    def style_workbook(ws, header_row, widths, title_rows=()):
        """Polished, readable Excel formatting with RTL layout and comfortable spacing."""
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.showGridLines = False
        navy = "17365D"
        teal = "DDEBF7"
        pale = "F4F8FC"
        white = "FFFFFF"
        line = Side(style="thin", color="D6E0EA")
        for r in title_rows:
            for cell in ws[r]:
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = Font(name="Aptos", size=12, bold=True, color=white)
                cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            ws.row_dimensions[r].height = 27
        for cell in ws[header_row]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(name="Aptos", size=11, bold=True, color=white)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(top=line, bottom=line, left=line, right=line)
        ws.row_dimensions[header_row].height = 32
        for row_idx in range(header_row + 1, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 30
            for cell in ws[row_idx]:
                cell.fill = PatternFill("solid", fgColor=white if row_idx % 2 else pale)
                cell.font = Font(name="Aptos", size=10, color="263445")
                cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
                cell.border = Border(top=line, bottom=line, left=line, right=line)
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = f"A{header_row + 1}"
        if ws.max_row >= header_row:
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.35
        ws.page_margins.right = 0.35
        ws.page_margins.top = 0.55
        ws.page_margins.bottom = 0.55
        ws.print_title_rows = f"{header_row}:{header_row}"

    async def choose_excel_path(default_name):
        """Cross-platform destination: Windows dialog; mobile-safe app export folder."""
        is_android = (os.environ.get("ANDROID_ARGUMENT") is not None or
                      os.environ.get("ANDROID_ROOT") is not None)
        if not is_android and sys.platform.startswith("win"):
            try:
                import tkinter as tk
                from tkinter import filedialog
                root = tk.Tk()
                root.withdraw()
                try:
                    root.attributes("-topmost", True)
                except Exception:
                    pass
                chosen = filedialog.asksaveasfilename(
                    parent=root,
                    title="انتخاب محل ذخیره فایل اکسل",
                    initialfile=default_name,
                    defaultextension=".xlsx",
                    filetypes=[("Excel Workbook", "*.xlsx")],
                )
                root.destroy()
                if not chosen:
                    notify("ذخیره فایل لغو شد.")
                    return None
                path = Path(chosen)
                return path if path.suffix.lower() == ".xlsx" else path.with_suffix(".xlsx")
            except Exception:
                # If the desktop dialog is unavailable, use the app's standard export folder.
                pass

        # Android and other non-Windows platforms save without opening tkinter/FilePicker.
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        path = EXPORT_DIR / default_name
        if path.suffix.lower() != ".xlsx":
            path = path.with_suffix(".xlsx")
        # Avoid silently overwriting a previous export with the same name.
        if path.exists():
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            path = path.with_name(f"{path.stem}_{stamp}{path.suffix}")
        return path

    def format_answer(item, value):
        if item["type"] == "checkbox":
            return item.get("label", "تیک") if value is True else "تیک نخورده" if value is False else "بدون پاسخ"
        if item["type"] == "yesno":
            return item.get("yes_label", "بله") if value == "yes" else item.get("no_label", "خیر") if value == "no" else "بدون پاسخ"
        return value if value not in (None, "") else "بدون پاسخ"

    async def export_response(rid):
        if Workbook is None:
            notify("برای خروجی اکسل اجرا کن: python -m pip install openpyxl")
            return
        with store.connect() as c:
            row = c.execute("SELECT * FROM responses WHERE id=?", (rid,)).fetchone()
        if not row:
            notify("پاسخ پیدا نشد.")
            return
        path = await choose_excel_path(f"پاسخ_{rid}.xlsx")
        if path is None:
            return
        try:
            payload = json.loads(row["answers"])
            cfg = normalize(json.dumps(payload.get("config", {}), ensure_ascii=False))
            answers = payload.get("answers", {})
            wb = Workbook()
            ws = wb.active
            ws.title = "پاسخ"
            ws.append(["عنوان پرسشنامه", row["questionnaire_title"]])
            ws.append(["تاریخ ثبت", row["created_at"]])
            ws.append(["بخش", "سؤال", "پاسخ"])
            for si, sec in enumerate(cfg["sections"]):
                section = answers.get(str(si), {})
                for qi, item in enumerate(sec["questions"]):
                    value = "بخش رد شد" if section.get("skipped") else format_answer(item, section.get("items", {}).get(str(qi)))
                    ws.append([sec["title"], item["prompt"], value])
            style_workbook(ws, 3, [28, 52, 36], title_rows=(1, 2))
            path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(path)
            notify(f"فایل اکسل ذخیره شد: {path}")
        except Exception as ex:
            notify(f"ساخت فایل اکسل ناموفق بود: {ex}")

    async def export_all(e=None):
        rows = store.responses()
        if not rows:
            notify("هنوز پاسخی ثبت نشده.")
            return
        if Workbook is None:
            notify("برای خروجی اکسل اجرا کن: python -m pip install openpyxl")
            return
        path = await choose_excel_path("همه_پاسخ‌ها.xlsx")
        if path is None:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "همه پاسخ‌ها"
            ws.append(["شناسه", "پرسشنامه", "تاریخ ثبت", "بخش", "سؤال", "پاسخ"])
            for row in rows:
                payload = json.loads(row["answers"])
                cfg = normalize(json.dumps(payload.get("config", {}), ensure_ascii=False))
                answers = payload.get("answers", {})
                for si, sec in enumerate(cfg["sections"]):
                    section = answers.get(str(si), {})
                    for qi, item in enumerate(sec["questions"]):
                        value = "بخش رد شد" if section.get("skipped") else format_answer(item, section.get("items", {}).get(str(qi)))
                        ws.append([row["id"], row["questionnaire_title"], row["created_at"], sec["title"], item["prompt"], value])
            style_workbook(ws, 1, [10, 30, 22, 28, 52, 36])
            path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(path)
            notify(f"فایل اکسل ذخیره شد: {path}")
        except Exception as ex:
            notify(f"ساخت فایل اکسل ناموفق بود: {ex}")

    def responses_view(qid=None):
        rows = store.responses(qid)
        controls = [header("پاسخ‌های ثبت‌شده", lambda e: home())]
        if not rows:
            controls.append(text("هنوز پاسخی ثبت نشده.", 15, gray))
        for row in rows:
            controls.append(ft.Card(content=ft.Container(padding=12, content=ft.Column([
                text(f'#{row["id"]} · {row["questionnaire_title"]}', 15, bold=True),
                text(row["created_at"], 12, gray),
                button("خروجی اکسل این پاسخ", lambda e, rid=row["id"]: page.run_task(export_response, rid), purple, ft.Icons.FILE_DOWNLOAD)
            ]))))
        controls.append(button("خروجی اکسل همه پاسخ‌ها", lambda e: page.run_task(export_all), green, ft.Icons.FILE_DOWNLOAD))
        page.controls.clear()
        page.add(ft.Column(controls, spacing=10, scroll=ft.ScrollMode.AUTO, expand=True))
        page.update()

    home()

if __name__ == "__main__":
    if hasattr(ft, "run"):
        ft.run(main)
    else:
        ft.app(target=main)
